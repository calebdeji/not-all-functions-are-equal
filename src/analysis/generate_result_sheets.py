"""
generate_result_sheets.py

Generates three Excel workbooks from the HiGraph no-overlap evaluation results:
  1. table1_pipeline_comparison.xlsx  — SDL variants vs baselines (FCGAT, FAMCF, GraphCL, SA, graph pooling)
  2. table2_sdl_variants.xlsx         — SDL-Degree vs SDL-Complexity vs SDL-PageRank
  3. table3_ablation.xlsx             — Backbone ablation (backbone-only + SDL-Complexity head)

Each file has two sheets:
  Accuracy   — 1-shot and 5-shot accuracy ± 95% CI
  Resources  — total params, head params, wall time, GPU time, peak memory
"""

import os
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

from src.core.utils import Utils

OUT_DIR = os.path.join(Utils.get_project_root(), 'results')
os.makedirs(OUT_DIR, exist_ok=True)

# ---------------------------------------------------------------------------
# Raw data — HiGraph, no-overlap, 5-way, 500 episodes
# ---------------------------------------------------------------------------

ACC = {
    # (1shot_mean, 1shot_ci, 5shot_mean, 5shot_ci)
    'sdl_complexity':           (66.46, 1.18, 80.35, 0.90),
    'sdl_degree':               (66.47, 1.18, 80.34, 0.90),
    'sdl_pagerank':             (66.45, 1.18, 80.34, 0.90),
    'self_attention':           (61.88, 1.15, 77.06, 0.94),
    'gmt':                      (63.01, 1.17, 74.70, 0.95),
    'gmpool':                   (64.69, 1.20, 77.54, 0.93),
    'topk_pool':                (59.17, 1.25, 73.89, 0.96),
    'fcgat':                    (59.78, 1.11, 77.07, 0.90),
    'famcf':                    (67.09, 1.18, 78.32, 0.94),
    'pretrained_backbone':      (57.67, 1.24, 71.61, 0.99),
    'androfim':                 (59.97, 1.20, 76.53, 0.97),
    # SDL ablation variants
    'sdl_no_learnable_param':   (66.35, 1.17, 79.93, 0.90),
    'sdl_complexity_prior':     (66.46, 1.16, 80.54, 0.90),
    'sdl_with_intra_param':     (66.47, 1.17, 80.34, 0.90),
    'sdl_without_structure':    (66.53, 1.18, 80.22, 0.91),
    # Backbone ablation
    'whole_graph_backbone':         (54.35, 1.18, 67.37, 1.01),
    'whole_graph_sdl_complexity':   (62.04, 1.18, 79.26, 0.87),
    'flat_fcg_backbone':            (53.50, 1.25, 67.72, 1.03),
    'flat_fcg_sdl_complexity':      (64.64, 1.15, 80.86, 0.84),
    'set_only_backbone':            (60.35, 1.31, 75.69, 0.97),
    'set_only_sdl_complexity':      (67.61, 1.23, 81.26, 0.91),
    'instance_only_backbone':       (54.19, 1.31, 66.72, 1.06),
    'instance_only_sdl_complexity': (66.63, 1.17, 82.71, 0.81),
}

PARAMS = {
    # (total_params, head_params)
    'sdl_complexity':           (1_584_898,   2),
    'sdl_degree':               (1_584_898,   2),
    'sdl_pagerank':             (1_584_898,   2),
    'self_attention':           (1_880_577,   295_681),
    'gmt':                      (9_284_864,   7_699_968),
    'gmpool':                   (4_553_216,   2_968_320),
    'topk_pool':                (2_176_256,   591_360),
    'fcgat':                    (532_264,     532_264),
    'famcf':                    (86_336,      86_336),   # sum of 3 encoders
    'pretrained_backbone':      (1_584_896,   1_584_896),
    'androfim':                 (7_996_800,   7_996_800),
    'sdl_no_learnable_param':   (1_584_896,   0),
    'sdl_complexity_prior':     (1_584_896,   0),
    'sdl_with_intra_param':     (1_584_897,   1),
    'sdl_without_structure':    (1_584_898,   2),
    'whole_graph_backbone':         (1_584_896, 1_584_896),
    'whole_graph_sdl_complexity':   (1_584_898, 2),
    'flat_fcg_backbone':            (1_316_608, 1_316_608),
    'flat_fcg_sdl_complexity':      (1_316_610, 2),
    'set_only_backbone':            (1_584_896, 1_584_896),
    'set_only_sdl_complexity':      (1_584_898, 2),
    'instance_only_backbone':       (1_584_896, 1_584_896),
    'instance_only_sdl_complexity': (1_584_898, 2),
}

# wall_mean, wall_std, gpu_mean, gpu_std, peak_mem_avg, peak_mem_max  (times in ms, mem in MB)
RESOURCES = {
    'sdl_complexity':           (29.51,  8.30,  29.56,  8.30,  1075.5, 2064.5),
    'sdl_degree':               (29.63,  8.34,  29.68,  8.34,  1075.5, 2064.5),
    'sdl_pagerank':             (34.64,  8.43,  34.69,  8.43,  1075.5, 2064.5),
    'self_attention':           (31.16, 11.34,  31.21, 11.34,  1075.5, 2064.5),
    'gmt':                      (54.19, 15.37,  54.24, 15.37,  3156.4, 6199.6),
    'gmpool':                   (52.52, 14.47,  52.57, 14.47,  3156.4, 6199.6),
    'topk_pool':                (29.80,  8.92,  29.84,  8.92,  1075.5, 2064.5),
    'fcgat':                    (14.04,  3.45,  14.08,  3.45,   613.5, 1086.6),
    'famcf':                    ( 2.69,  0.63,   2.73,  0.63,   192.0,  218.4),
    'pretrained_backbone':      (24.64,  7.17,  24.68,  7.17,  1037.8, 1952.9),
    'androfim':                 (55348.11, 35965.27, 55347.90, 35965.06, 9800.9, 16236.1),
    'sdl_no_learnable_param':   (29.14,  8.34,  29.18,  8.34,  1075.5, 2064.5),
    'sdl_complexity_prior':     (28.44,  8.12,  28.48,  8.12,  1075.5, 2064.5),
    'sdl_with_intra_param':     (29.26,  8.34,  29.30,  8.34,  1075.5, 2064.5),
    'sdl_without_structure':    (29.24,  8.36,  29.29,  8.36,  1075.5, 2064.5),
    'whole_graph_backbone':         (24.46,  7.02,  24.50,  7.02,  1037.8, 1952.9),
    'whole_graph_sdl_complexity':   (27.39,  7.61,  27.43,  7.61,  1075.5, 2064.5),
    'flat_fcg_backbone':            (10.57,  2.89,  10.60,  2.89,   534.9,  863.6),
    'flat_fcg_sdl_complexity':      (13.15,  3.36,  13.18,  3.36,   572.6,  968.7),
    'set_only_backbone':            (24.84,  7.22,  24.88,  7.22,  1037.8, 1952.9),
    'set_only_sdl_complexity':      (27.33,  7.62,  27.37,  7.62,  1075.5, 2064.5),
    'instance_only_backbone':       (27.28,  7.99,  27.32,  7.99,  1037.8, 1952.9),
    'instance_only_sdl_complexity': (29.72,  8.37,  29.76,  8.37,  1075.5, 2064.5),
}

DISPLAY_NAMES = {
    'sdl_complexity':               'SDL-Complexity (Ours)',
    'sdl_degree':                   'SDL-Degree (Ours)',
    'sdl_pagerank':                 'SDL-PageRank (Ours)',
    'self_attention':               'Self-Attention',
    'gmt':                          'GMT',
    'gmpool':                       'GMPool',
    'topk_pool':                    'TopK-Pool',
    'fcgat':                        'FCGAT',
    'famcf':                        'FAMCF',
    'pretrained_backbone':          'GraphCL Backbone',
    'androfim':                     'AndroFIM',
    'sdl_no_learnable_param':       'SDL (No Learnable Param)',
    'sdl_complexity_prior':         'SDL-Complexity (Prior)',
    'sdl_with_intra_param':         'SDL (With Intra-Param)',
    'sdl_without_structure':        'SDL (Without Structure)',
    'whole_graph_backbone':         'Whole-Graph Backbone',
    'whole_graph_sdl_complexity':   'Whole-Graph + SDL-Complexity',
    'flat_fcg_backbone':            'Flat-FCG Backbone',
    'flat_fcg_sdl_complexity':      'Flat-FCG + SDL-Complexity',
    'set_only_backbone':            'Set-Only Backbone',
    'set_only_sdl_complexity':      'Set-Only + SDL-Complexity',
    'instance_only_backbone':       'Instance-Only Backbone',
    'instance_only_sdl_complexity': 'Instance-Only + SDL-Complexity',
}

# ---------------------------------------------------------------------------
# Styling helpers
# ---------------------------------------------------------------------------

HEADER_FILL   = PatternFill('solid', fgColor='1F4E79')
SUBHDR_FILL   = PatternFill('solid', fgColor='2E75B6')
OUR_FILL      = PatternFill('solid', fgColor='E2EFDA')   # light green for SDL rows
ALT_FILL      = PatternFill('solid', fgColor='F2F2F2')
WHITE_FILL    = PatternFill('solid', fgColor='FFFFFF')

HEADER_FONT   = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
SUBHDR_FONT   = Font(name='Calibri', bold=True, color='FFFFFF', size=10)
BODY_FONT     = Font(name='Calibri', size=10)
BOLD_FONT     = Font(name='Calibri', bold=True, size=10)

CENTER  = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT    = Alignment(horizontal='left',   vertical='center')
RIGHT   = Alignment(horizontal='right',  vertical='center')

THIN = Side(style='thin', color='BFBFBF')
MED  = Side(style='medium', color='595959')

def thin_border():
    return Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def med_border():
    return Border(left=MED, right=MED, top=MED, bottom=MED)

def set_col_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def write_header_cell(ws, row, col, value, fill=None, font=None, align=CENTER):
    c = ws.cell(row=row, column=col, value=value)
    c.fill   = fill  or HEADER_FILL
    c.font   = font  or HEADER_FONT
    c.alignment = align
    c.border = thin_border()
    return c

def write_body_cell(ws, row, col, value, fill=None, font=None, align=None, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.fill   = fill or WHITE_FILL
    c.font   = font or BODY_FONT
    c.alignment = align or CENTER
    c.border = thin_border()
    if fmt:
        c.number_format = fmt
    return c

def freeze_and_zoom(ws, freeze='B3'):
    ws.freeze_panes = freeze
    ws.sheet_view.zoomScale = 100

# ---------------------------------------------------------------------------
# Accuracy sheet builder
# ---------------------------------------------------------------------------

def build_accuracy_sheet(ws, rows, title, note=''):
    """
    rows: list of (key, section_label_or_None)
    section_label_or_None: if not None, insert a section divider row with that label
    """
    ws.title = 'Accuracy'

    # Title row
    ws.merge_cells('A1:H1')
    tc = ws['A1']
    tc.value     = title
    tc.font      = Font(name='Calibri', bold=True, size=13, color='1F4E79')
    tc.alignment = LEFT
    if note:
        ws.merge_cells('A2:H2')
        nc = ws['A2']
        nc.value     = note
        nc.font      = Font(name='Calibri', italic=True, size=9, color='595959')
        nc.alignment = LEFT

    hdr_row = 3
    # Headers
    headers = ['Model', '1-Shot Acc (%)', '1-Shot ±CI', '5-Shot Acc (%)', '5-Shot ±CI',
               'Δ Acc (5-1)', 'Episodes', 'Dataset']
    for ci, h in enumerate(headers, 1):
        write_header_cell(ws, hdr_row, ci, h)

    data_start = hdr_row + 1
    row_idx = data_start

    for item in rows:
        if isinstance(item, str):
            # Section header
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=8)
            c = ws.cell(row=row_idx, column=1, value=item)
            c.fill      = SUBHDR_FILL
            c.font      = SUBHDR_FONT
            c.alignment = LEFT
            c.border    = thin_border()
            row_idx += 1
            continue

        key = item
        name = DISPLAY_NAMES.get(key, key)
        a1m, a1c, a5m, a5c = ACC[key]
        delta = a5m - a1m
        is_ours = key.startswith('sdl_')
        fill = OUR_FILL if is_ours else (ALT_FILL if row_idx % 2 == 0 else WHITE_FILL)

        vals = [name, a1m, a1c, a5m, a5c, delta, 500, 'HiGraph (no-overlap)']
        aligns = [LEFT, CENTER, CENTER, CENTER, CENTER, CENTER, CENTER, CENTER]
        fmts   = [None, '0.00', '±0.00', '0.00', '±0.00', '+0.00;-0.00', '0', None]
        for ci, (v, al, fm) in enumerate(zip(vals, aligns, fmts), 1):
            write_body_cell(ws, row_idx, ci, v, fill=fill, align=al, fmt=fm)

        row_idx += 1

    set_col_widths(ws, {
        'A': 34, 'B': 14, 'C': 10, 'D': 14, 'E': 10, 'F': 12, 'G': 10, 'H': 22,
    })
    freeze_and_zoom(ws, 'B4')
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 14
    ws.row_dimensions[3].height = 32


# ---------------------------------------------------------------------------
# Resource sheet builder
# ---------------------------------------------------------------------------

def build_resource_sheet(ws, keys, title):
    ws.title = 'Resources'

    ws.merge_cells('A1:I1')
    tc = ws['A1']
    tc.value     = title + ' — Resource Usage'
    tc.font      = Font(name='Calibri', bold=True, size=13, color='1F4E79')
    tc.alignment = LEFT

    hdr_row = 3
    headers = [
        'Model',
        'Total Params',
        'Head Params',
        'Wall Time (ms)',
        'Wall ±Std',
        'GPU Time (ms)',
        'GPU ±Std',
        'Peak Mem Avg (MB)',
        'Peak Mem Max (MB)',
    ]
    for ci, h in enumerate(headers, 1):
        write_header_cell(ws, hdr_row, ci, h)

    row_idx = hdr_row + 1
    for i, key in enumerate(keys):
        if isinstance(key, str) and key not in ACC:
            # Section label
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=9)
            c = ws.cell(row=row_idx, column=1, value=key)
            c.fill      = SUBHDR_FILL
            c.font      = SUBHDR_FONT
            c.alignment = LEFT
            c.border    = thin_border()
            row_idx += 1
            continue

        name = DISPLAY_NAMES.get(key, key)
        tp, hp = PARAMS.get(key, (0, 0))
        wm, ws_, gm, gs_, pm, pmax = RESOURCES.get(key, (0,)*6)
        is_ours = key.startswith('sdl_')
        fill = OUR_FILL if is_ours else (ALT_FILL if row_idx % 2 == 0 else WHITE_FILL)

        vals   = [name, tp, hp, wm, ws_, gm, gs_, pm, pmax]
        aligns = [LEFT] + [RIGHT]*8
        fmts   = [None, '#,##0', '#,##0', '0.00', '0.00', '0.00', '0.00', '0.0', '0.0']
        for ci, (v, al, fm) in enumerate(zip(vals, aligns, fmts), 1):
            write_body_cell(ws, row_idx, ci, v, fill=fill, align=al, fmt=fm)

        row_idx += 1

    set_col_widths(ws, {
        'A': 34, 'B': 14, 'C': 13, 'D': 14, 'E': 10,
        'F': 14, 'G': 10, 'H': 18, 'I': 18,
    })
    freeze_and_zoom(ws, 'B4')
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[3].height = 32


# ---------------------------------------------------------------------------
# Table 1 — Pipeline comparison
# ---------------------------------------------------------------------------

def make_table1():
    wb = Workbook()
    ws_acc = wb.active
    ws_res = wb.create_sheet()

    T1_TITLE = 'Table 1: Pipeline Comparison — HiGraph No-Overlap (5-way, 500 episodes)'
    T1_NOTE  = 'SDL variants highlighted in green. Random-chance baseline = 20.00%.'

    rows = [
        'SDL Variants (Ours)',
        'sdl_complexity',
        'sdl_degree',
        'sdl_pagerank',
        'Graph Pooling Baselines',
        'gmt',
        'gmpool',
        'topk_pool',
        'Other Baselines',
        'self_attention',
        'fcgat',
        'famcf',
        'pretrained_backbone',
        'androfim',
    ]

    build_accuracy_sheet(ws_acc, rows, T1_TITLE, T1_NOTE)

    flat_keys = [r for r in rows if r in ACC]
    build_resource_sheet(ws_res, rows, T1_TITLE)

    path = os.path.join(OUT_DIR, 'table1_pipeline_comparison.xlsx')
    wb.save(path)
    print(f'Saved: {path}')


# ---------------------------------------------------------------------------
# Table 2 — SDL variants
# ---------------------------------------------------------------------------

def make_table2():
    wb = Workbook()
    ws_acc = wb.active
    ws_res = wb.create_sheet()

    T2_TITLE = 'Table 2: SDL Variants — HiGraph No-Overlap (5-way, 500 episodes)'
    T2_NOTE  = 'Comparison of structural difficulty signal: degree, cyclomatic complexity, PageRank.'

    rows = [
        'sdl_degree',
        'sdl_complexity',
        'sdl_pagerank',
    ]

    build_accuracy_sheet(ws_acc, rows, T2_TITLE, T2_NOTE)
    build_resource_sheet(ws_res, rows, T2_TITLE)

    path = os.path.join(OUT_DIR, 'table2_sdl_variants.xlsx')
    wb.save(path)
    print(f'Saved: {path}')


# ---------------------------------------------------------------------------
# Table 3 — Ablation
# ---------------------------------------------------------------------------

def make_table3():
    wb = Workbook()
    ws_acc = wb.active
    ws_res = wb.create_sheet()

    T3_TITLE = 'Table 3: Ablation Study — HiGraph No-Overlap (5-way, 500 episodes)'
    T3_NOTE  = ('Backbone rows = no episodic head (mean-pool at inference). '
                'SDL-Complexity rows = episodic head trained on same backbone.')

    rows = [
        'Full Backbone (Instance + Set NT-Xent)',
        'pretrained_backbone',
        'sdl_complexity',
        'Whole-Graph Backbone (no hierarchical decomposition)',
        'whole_graph_backbone',
        'whole_graph_sdl_complexity',
        'Flat-FCG Backbone (no CFG)',
        'flat_fcg_backbone',
        'flat_fcg_sdl_complexity',
        'Set-Only Backbone (set-level NT-Xent only)',
        'set_only_backbone',
        'set_only_sdl_complexity',
        'Instance-Only Backbone (instance-level NT-Xent only)',
        'instance_only_backbone',
        'instance_only_sdl_complexity',
    ]

    build_accuracy_sheet(ws_acc, rows, T3_TITLE, T3_NOTE)
    build_resource_sheet(ws_res, rows, T3_TITLE)

    path = os.path.join(OUT_DIR, 'table3_ablation.xlsx')
    wb.save(path)
    print(f'Saved: {path}')


# ---------------------------------------------------------------------------

if __name__ == '__main__':
    make_table1()
    make_table2()
    make_table3()
    print('Done.')
